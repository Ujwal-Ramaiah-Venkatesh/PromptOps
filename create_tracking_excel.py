#!/usr/bin/env python3
"""
Create PromptOps_Tracking.xlsx with all 8 tracking sheets
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import csv
import os

def create_excel_from_csvs():
    """Create single Excel workbook from all CSV files"""

    # Create new workbook
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default sheet

    # Define sheets with their names and CSV files
    sheets = [
        ("Daily Work Log", "tracking/01_daily_work_log.csv"),
        ("Task Tracker", "tracking/02_task_tracker.csv"),
        ("Deliverables", "tracking/03_deliverables_tracker.csv"),
        ("Time Tracking", "tracking/04_time_tracking.csv"),
        ("Issues & Blockers", "tracking/05_issues_blockers.csv"),
        ("Weekly Summary", "tracking/06_weekly_summary.csv"),
        ("Exit Criteria", "tracking/07_phase1_exit_criteria.csv"),
        ("Team Roster", "tracking/08_team_roster.csv"),
    ]

    # Header style
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for sheet_name, csv_file in sheets:
        print(f"Creating sheet: {sheet_name}")

        # Create new sheet
        ws = wb.create_sheet(title=sheet_name)

        # Read CSV and write to sheet
        with open(csv_file, 'r', encoding='utf-8') as f:
            reader = csv.reader(f)
            for row_idx, row in enumerate(reader, start=1):
                for col_idx, value in enumerate(row, start=1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=value)

                    # Style header row
                    if row_idx == 1:
                        cell.font = header_font
                        cell.fill = header_fill
                        cell.alignment = header_alignment

        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)

            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass

            adjusted_width = min(max_length + 2, 50)  # Cap at 50
            ws.column_dimensions[column_letter].width = adjusted_width

        # Freeze header row
        ws.freeze_panes = "A2"

    # Save workbook
    output_file = "PromptOps_Tracking.xlsx"
    wb.save(output_file)
    print(f"\n✅ Created: {output_file}")
    print(f"Location: {os.path.abspath(output_file)}")
    print(f"Sheets: {len(sheets)}")

    return output_file

if __name__ == "__main__":
    create_excel_from_csvs()
